# -*- coding: utf-8 -*-
"""按成都文理学院《软件系统类论文模板》格式生成毕业论文 Word 文档。

格式要点：
- 主标题：黑体二号；副标题黑体小二
- 一级标题(章)：黑体，三号(16pt)加粗；每章分页
- 二级标题：黑体，四号(14pt)
- 三级标题：黑体，小四(12pt)
- 正文：宋体小四(12pt)，行距固定 20 磅(254000 EMU)，首行缩进 2 字符
- 图题：黑体五号(10.5pt)，图题在图下方居中
- 表题：黑体五号(10.5pt)，表题在表上方居中
- 英文用 Times New Roman
"""
from docx import Document
from docx.shared import Pt, Emu, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = "D:/Demo/论文/课程助教系统毕业论文.docx"


def set_run(run, cn_font="宋体", en_font="Times New Roman", size=12, bold=False):
    run.font.name = en_font
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), en_font)
    rFonts.set(qn('w:hAnsi'), en_font)
    rFonts.set(qn('w:eastAsia'), cn_font)


def add_para(doc, text, cn_font="宋体", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             line=20, first_indent=True):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Emu(int(line * 12700))
    if first_indent:
        pf.first_line_indent = Pt(2 * size)
    r = p.add_run(text)
    set_run(r, cn_font=cn_font, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1, size=16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Emu(int(20 * 12700))
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    if level == 1:
        p.paragraph_format.page_break_before = True
    r = p.add_run(text)
    set_run(r, cn_font="黑体", size=size, bold=True)
    return p


def add_caption(doc, text, above=False):
    """图表题注：黑体五号(10.5pt)，居中。表题在上方，图题在下方调用。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Emu(int(20 * 12700))
    r = p.add_run(text)
    set_run(r, cn_font="黑体", size=10.5, bold=False)
    return p


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run(r, cn_font="宋体", size=10.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            set_run(r, cn_font="宋体", size=10.5, bold=False)
    return t


def add_placeholder(doc, text):
    add_para(doc, "【" + text + "】", cn_font="宋体", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)


doc = Document()
# 默认字体
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# ================= 封 面 =================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("本科毕业论文（设计）"); set_run(r, cn_font="黑体", size=22, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("基于 Agent 的课程助教系统"); set_run(r, cn_font="黑体", size=22, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("的设计与实现"); set_run(r, cn_font="黑体", size=22, bold=True)
for _ in range(3):
    doc.add_paragraph()
for line in ["学    院：人工智能与大数据学院", "专    业：计算机科学与技术",
             "学    号：＿＿＿＿＿＿＿＿", "姓    名：＿＿＿＿＿＿＿＿"]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(line); set_run(r, cn_font="宋体", size=15)
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("二〇二六年四月"); set_run(r, cn_font="宋体", size=15)
doc.add_page_break()


def read_md(path, start=0):
    with open(path, encoding="utf-8") as f:
        lines = f.readlines()
    return lines[start:]


def render(md_path, skip_title=0):
    lines = read_md(md_path)
    for ln in lines[skip_title:]:
        s = ln.rstrip("\n")
        if not s.strip():
            continue
        if s.startswith("# "):  # 章标题（一级）
            add_heading(doc, s[2:].strip(), 1, 16)
        elif s.startswith("## "):
            add_heading(doc, s[3:].strip(), 2, 14)
        elif s.startswith("### "):
            add_heading(doc, s[4:].strip(), 3, 12)
        elif s.startswith("```"):
            continue
        elif s.strip().startswith("【"):
            add_placeholder(doc, s.strip().strip("【】"))
            continue
        elif (s[:1].startswith(("表", "图"))
              and len(s) > 5
              and s[1:2].isdigit()
              and " " in s
              and s.split(" ", 1)[1].strip()):
            # 图表题注（如"表4.1 用户表结构"、"图5.2 学生答疑界面"）——黑体五号居中
            add_caption(doc, s.strip())
            continue
        else:
            # 参考文献条目 / 致谢空段等普通段落
            add_para(doc, s, cn_font="宋体", size=12)


# 中文摘要（第1行是 "## 摘  要" 之类）
# 我们手动渲染摘要/Abstract 部分
render_abs = True

# 摘要
add_para(doc, "摘  要", cn_font="黑体", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
abs_text = (
    "随着人工智能技术的快速发展，大语言模型在自然语言处理领域取得了显著进展，并在教育领域展现出广阔的应用前景。"
    "传统高校课程助教虽能完成答疑、作业批改与出题等工作，但普遍存在响应不及时、批改标准不统一、依赖人工重复劳动等问题，"
    "难以满足大规模个性化教学的需求。针对上述问题，本文设计并实现了一个基于 Agent 架构的课程助教系统。"
    "系统以 LangChain 与 LangGraph 为核心框架，构建了答疑、批改、出题三个相互协作的智能体：答疑 Agent 结合检索增强生成（RAG）技术，"
    "基于课程知识库进行流式回答并附带溯源引用；批改 Agent 依据预设评分标准对学生作业进行智能评分并生成评语；"
    "出题 Agent 依据知识点与难度批量生成试题。系统后端基于 FastAPI 构建 Web 服务并接入云端大模型 API，"
    "前端基于 Jinja2 模板实现学生、教师、管理员三端页面，实现流式对话、知识库管理、作业批改、试题管理与学情看板等核心功能。"
    "系统通过关键词路由实现多 Agent 协作调度，并利用 Function Calling 机制完成函数调用闭环。"
    "本文详细阐述了系统的需求分析、总体设计、详细设计、数据库设计及系统实现，并通过功能测试验证了各模块的可用性与正确性。"
    "测试结果表明，系统能够有效支持课程教学中答疑、批改与出题等核心环节，具有良好的可用性与扩展性。"
)
add_para(doc, abs_text, cn_font="宋体", size=12)
add_para(doc, "关键词：大语言模型；检索增强生成；多智能体；答疑系统；课程助教", cn_font="宋体", size=12, first_indent=False)
doc.add_page_break()

# 英文摘要
add_para(doc, "Abstract", cn_font="Times New Roman", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
abs_en = (
    "With the rapid development of artificial intelligence, large language models (LLMs) have made significant progress "
    "in the field of natural language processing and shown broad application prospects in education. Traditional university "
    "course teaching assistants can complete tasks such as answering questions, grading assignments, and generating exercises, "
    "but generally suffer from problems such as untimely responses, inconsistent grading standards, and heavy reliance on "
    "repetitive manual labor. To address these problems, this thesis designs and implements a course teaching assistant system "
    "based on the Agent architecture. Using LangChain and LangGraph as the core frameworks, the system builds three collaborating "
    "agents: a Q&A agent, a grading agent, and a question-generation agent. The Q&A agent combines Retrieval-Augmented Generation "
    "(RAG) to provide streaming answers based on the course knowledge base with source tracing. The grading agent scores assignments "
    "according to preset criteria and generates feedback. The question-generation agent generates exercises in batches according to "
    "knowledge points and difficulty. The back end is built with FastAPI and connects to cloud-based LLM APIs, while the front end "
    "uses Jinja2 templates for student, teacher, and administrator pages. The system implements multi-agent scheduling through "
    "keyword-based routing and completes the function-calling loop through Function Calling. This thesis elaborates on the system's "
    "requirements analysis, overall design, detailed design, database design, and implementation, and verifies the usability and "
    "correctness of each module through functional testing. Test results show that the system can effectively support the core links "
    "of Q&A, grading, and question generation in course teaching."
)
add_para(doc, abs_en, cn_font="宋体", size=12)
add_para(doc, "Key words: Large Language Model; Retrieval-Augmented Generation; Multi-Agent; Q&A System; Course Teaching Assistant",
         cn_font="Times New Roman", size=12, first_indent=False)
doc.add_page_break()

# 目录占位
add_para(doc, "目  录", cn_font="黑体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
for t in ["1　绪论", "2　系统关键技术介绍", "3　系统需求分析", "4　系统设计",
          "5　系统实现", "6　总结与展望", "参考文献", "致  谢"]:
    add_para(doc, t, cn_font="宋体", size=12, first_indent=False)
doc.add_page_break()

# 正文：第1章起 每章分页（标题 page_break_before）
# 手动渲染第1-6章正文，合并到一个文件最简单：直接逐个渲染
render("D:/Demo/论文/thesis_ch1_2.md", skip_title=0)
render("D:/Demo/论文/thesis_ch2_keytech.md", skip_title=0)
render("D:/Demo/论文/thesis_ch3_requirement.md", skip_title=0)
render("D:/Demo/论文/thesis_ch4_design.md", skip_title=0)
render("D:/Demo/论文/thesis_ch5_impl.md", skip_title=0)
render("D:/Demo/论文/thesis_ch6_conclusion.md", skip_title=0)

doc.save(OUT)
print("saved:", OUT)